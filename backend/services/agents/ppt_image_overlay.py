"""Legacy image-overlay PPTX export helpers.

The competition/direct-SVG pipeline is the default export path.  These
helpers remain isolated for explicitly requested legacy conversion modes:
image-only decks, bitmap backgrounds with editable text overlays, and native
outline drafts.  Keeping them out of ``ppt_agent`` makes the orchestrator
easier to reason about and prevents legacy implementation detail from leaking
into the direct-SVG route.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path


logger = logging.getLogger(__name__)


def build_pptx_from_slides(
    job_dir: Path,
    outline: dict,
    slide_images: list[bytes],
    output_path: Path,
) -> Path:
    """Build an editable PPTX from image pages and outline copy.

    This compatibility exporter uses a full-page image layer with native text
    boxes above it.  The direct-SVG exporter is preferred for new decks.
    """
    del job_dir  # Kept in the signature for callers of the legacy API.
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Emu(9144000)
    presentation.slide_height = Emu(5143500)
    blank_layout = presentation.slide_layouts[6]

    slides_info = outline.get("slides", [])
    for index, (image_bytes, slide_info) in enumerate(zip(slide_images, slides_info)):
        slide = presentation.slides.add_slide(blank_layout)
        picture = slide.shapes.add_picture(
            io.BytesIO(image_bytes),
            left=Emu(0),
            top=Emu(0),
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        _move_shape_to_back(slide, picture)

        title_box = slide.shapes.add_textbox(
            left=Inches(0.5), top=Inches(0.3), width=Inches(9), height=Inches(0.8)
        )
        title_frame = title_box.text_frame
        title_frame.word_wrap = True
        title_paragraph = title_frame.paragraphs[0]
        title_paragraph.alignment = PP_ALIGN.LEFT
        title_run = title_paragraph.add_run()
        title_run.text = slide_info.get("title", f"第 {index + 1} 页")
        title_run.font.size = Pt(28)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        points = slide_info.get("points", [])
        if points:
            content_box = slide.shapes.add_textbox(
                left=Inches(0.5), top=Inches(1.2), width=Inches(8), height=Inches(3.5)
            )
            content_frame = content_box.text_frame
            content_frame.word_wrap = True
            for point_index, point in enumerate(points):
                paragraph = content_frame.paragraphs[0] if point_index == 0 else content_frame.add_paragraph()
                paragraph.alignment = PP_ALIGN.LEFT
                run = paragraph.add_run()
                run.text = f"• {point}"
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        number_box = slide.shapes.add_textbox(
            left=Inches(9), top=Inches(5.2), width=Inches(0.8), height=Inches(0.3)
        )
        number_paragraph = number_box.text_frame.paragraphs[0]
        number_paragraph.alignment = PP_ALIGN.RIGHT
        number_run = number_paragraph.add_run()
        number_run.text = str(index + 1)
        number_run.font.size = Pt(10)
        number_run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    presentation.save(str(output_path))
    logger.info("[PPT image overlay] PPTX saved: %s", output_path)
    return output_path


def estimate_text_elements(slide_info: dict) -> list[dict]:
    """Estimate text geometry when legacy OCR is not available."""
    elements: list[dict] = []
    slide_type = slide_info.get("type", "content")
    title = slide_info.get("title", "")
    if title:
        if slide_type == "cover":
            elements.append({
                "content": title, "x": 0.05, "y": 0.35, "w": 0.90, "h": 0.15,
                "font_size_pt": 44, "bold": True, "color_hex": "#FFFFFF", "align": "center",
            })
        else:
            elements.append({
                "content": title, "x": 0.05, "y": 0.06, "w": 0.90, "h": 0.12,
                "font_size_pt": 32, "bold": True, "color_hex": "#FFFFFF", "align": "left",
            })

    if slide_type != "cover":
        for index, point in enumerate(slide_info.get("points", [])):
            y = 0.25 + index * 0.12
            if y + 0.10 > 0.95:
                break
            elements.append({
                "content": f"• {point}", "x": 0.06, "y": y, "w": 0.88, "h": 0.10,
                "font_size_pt": 20, "bold": False, "color_hex": "#EEEEEE", "align": "left",
            })
    return elements


def build_text_mask(image_bytes: bytes, text_elements: list[dict]) -> bytes:
    """Return a softly padded mask for text removal by an inpainting provider."""
    from PIL import Image, ImageDraw, ImageFilter

    image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    width, height = image.size
    mask = Image.new("L", (width, height), 0)
    draw = ImageDraw.Draw(mask)

    for item in text_elements:
        if not isinstance(item, dict) or not str(item.get("content", "")).strip():
            continue
        try:
            x, y, w, h = (float(item.get(name, 0)) for name in ("x", "y", "w", "h"))
        except (TypeError, ValueError):
            continue
        left, top = int(x * width), int(y * height)
        right, bottom = int((x + w) * width), int((y + h) * height)
        pad_x = max(8, int((right - left) * 0.16))
        pad_y = max(6, int((bottom - top) * 0.24))
        draw.rounded_rectangle(
            (
                max(0, left - pad_x), max(0, top - pad_y),
                min(width, right + pad_x), min(height, bottom + pad_y),
            ),
            radius=max(8, int(min(width, height) * 0.012)),
            fill=255,
        )

    output = io.BytesIO()
    mask.filter(ImageFilter.GaussianBlur(radius=2)).save(output, format="PNG")
    return output.getvalue()


def remove_text_regions_from_slide(image_bytes: bytes, text_elements: list[dict]) -> bytes:
    """Use a local colour-sampling cleanup if remote inpainting is unavailable."""
    if not text_elements:
        return image_bytes
    try:
        from PIL import Image, ImageDraw, ImageFilter

        image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
        width, height = image.size
        draw = ImageDraw.Draw(image, "RGBA")

        def clamp(value: int, low: int, high: int) -> int:
            return max(low, min(high, value))

        def fill_for(left: int, top: int, right: int, bottom: int) -> tuple[int, int, int, int]:
            pad = max(8, int(min(width, height) * 0.012))
            crop = image.crop((
                clamp(left - pad, 0, width), clamp(top - pad, 0, height),
                clamp(right + pad, 0, width), clamp(bottom + pad, 0, height),
            ))
            crop_width, crop_height = crop.size
            if crop_width <= 0 or crop_height <= 0:
                return (255, 255, 255, 238)
            pixels = []
            for x_pos in range(0, crop_width, max(1, crop_width // 16)):
                pixels.extend((crop.getpixel((x_pos, 0)), crop.getpixel((x_pos, crop_height - 1))))
            for y_pos in range(0, crop_height, max(1, crop_height // 16)):
                pixels.extend((crop.getpixel((0, y_pos)), crop.getpixel((crop_width - 1, y_pos))))
            if not pixels:
                return (255, 255, 255, 238)
            return (*[int(sum(pixel[channel] for pixel in pixels) / len(pixels)) for channel in range(3)], 238)

        for item in text_elements:
            if not isinstance(item, dict) or not str(item.get("content", "")).strip():
                continue
            try:
                x, y, w, h = (float(item.get(name, 0)) for name in ("x", "y", "w", "h"))
            except (TypeError, ValueError):
                continue
            left, top = int(x * width), int(y * height)
            right, bottom = int((x + w) * width), int((y + h) * height)
            pad_x = max(6, int((right - left) * 0.10))
            pad_y = max(4, int((bottom - top) * 0.18))
            left, top = clamp(left - pad_x, 0, width), clamp(top - pad_y, 0, height)
            right, bottom = clamp(right + pad_x, 0, width), clamp(bottom + pad_y, 0, height)
            if right <= left or bottom <= top:
                continue
            draw.rounded_rectangle(
                (left, top, right, bottom),
                radius=max(8, min(28, int(min(right - left, bottom - top) * 0.18))),
                fill=fill_for(left, top, right, bottom),
            )

        output = io.BytesIO()
        image.filter(ImageFilter.SMOOTH_MORE).save(output, format="PNG")
        return output.getvalue()
    except Exception as exc:
        logger.warning("[PPT image overlay] local text cleanup failed: %s", exc)
        return image_bytes


def build_pptx_with_text_overlay(
    slide_images: list[bytes],
    all_text_elements: list[list[dict]],
    slides_info: list[dict],
    output_path: Path,
    clean_backgrounds: list[bytes] | None = None,
) -> None:
    """Build a bitmap-background PPTX with editable text boxes above it."""
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    slide_width, slide_height = 9144000, 5143500
    presentation = Presentation()
    presentation.slide_width = Emu(slide_width)
    presentation.slide_height = Emu(slide_height)
    blank_layout = presentation.slide_layouts[6]
    alignments = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

    for index, image_bytes in enumerate(slide_images):
        slide_info = slides_info[index] if index < len(slides_info) else {}
        text_elements = all_text_elements[index] if index < len(all_text_elements) else estimate_text_elements(slide_info)
        if not isinstance(text_elements, list):
            text_elements = estimate_text_elements(slide_info)
        background = clean_backgrounds[index] if clean_backgrounds and index < len(clean_backgrounds) else b""
        background = background or remove_text_regions_from_slide(image_bytes, text_elements)
        slide = presentation.slides.add_slide(blank_layout)
        picture = slide.shapes.add_picture(
            io.BytesIO(background), left=Emu(0), top=Emu(0),
            width=Emu(slide_width), height=Emu(slide_height),
        )
        _move_shape_to_back(slide, picture)

        for element in text_elements:
            content = str(element.get("content", "")).strip()
            if not content:
                continue
            text_box = slide.shapes.add_textbox(
                left=Emu(int(element["x"] * slide_width)),
                top=Emu(int(element["y"] * slide_height)),
                width=Emu(int(element["w"] * slide_width)),
                height=Emu(int(element["h"] * slide_height)),
            )
            text_frame = text_box.text_frame
            text_frame.word_wrap = True
            text_frame.margin_left = text_frame.margin_right = Emu(0)
            text_frame.margin_top = text_frame.margin_bottom = Emu(0)
            text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            shape_properties = text_box._element.find(qn("p:spPr"))
            if shape_properties is not None:
                for fill_name in ("a:solidFill", "a:gradFill", "a:pattFill", "a:noFill"):
                    for fill in shape_properties.findall(qn(fill_name)):
                        shape_properties.remove(fill)
                etree.SubElement(shape_properties, qn("a:noFill"))

            paragraph = text_frame.paragraphs[0]
            paragraph.alignment = alignments.get(element.get("align", "left"), PP_ALIGN.LEFT)
            run = paragraph.add_run()
            run.text = content
            run.font.size = Pt(int(element.get("font_size_pt", 20)))
            run.font.bold = bool(element.get("bold", False))
            color = str(element.get("color_hex", "#FFFFFF")).lstrip("#")
            try:
                run.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
            except (TypeError, ValueError):
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    presentation.save(str(output_path))
    logger.info("[PPT image overlay] editable-overlay PPTX saved: %s", output_path)


def build_pptx_images_only(slide_images: list[bytes], output_path: Path) -> None:
    """Build a PPTX containing one full-page image per slide."""
    from pptx import Presentation
    from pptx.util import Emu

    slide_width, slide_height = 9144000, 5143500
    presentation = Presentation()
    presentation.slide_width = Emu(slide_width)
    presentation.slide_height = Emu(slide_height)
    blank_layout = presentation.slide_layouts[6]
    for image_bytes in slide_images:
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            io.BytesIO(image_bytes), left=Emu(0), top=Emu(0),
            width=Emu(slide_width), height=Emu(slide_height),
        )
    presentation.save(str(output_path))
    logger.info("[PPT image overlay] image-only PPTX saved: %s", output_path)


def build_pptx_from_outline(outline: dict, output_path: Path) -> None:
    """Build a simple native-outline draft for the explicitly requested legacy mode."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    title_color = RGBColor(0x18, 0x2A, 0x45)
    accent_color = RGBColor(0xFC, 0xA3, 0x11)
    body_color = RGBColor(0x35, 0x3F, 0x4F)
    muted_color = RGBColor(0x8A, 0x94, 0xA6)
    slides = outline.get("slides", []) or []
    deck_title = outline.get("title", "PPT")

    for index, info in enumerate(slides):
        slide = presentation.slides.add_slide(blank_layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(0xFA, 0xF7, 0xEF)
        band = slide.shapes.add_shape(1, Inches(0), Inches(0), presentation.slide_width, Inches(0.18))
        band.fill.solid()
        band.fill.fore_color.rgb = accent_color
        band.line.fill.background()
        slide_type = info.get("type", "content")
        title = info.get("title") or (deck_title if index == 0 else f"第 {index + 1} 页")

        if index == 0 or slide_type == "cover":
            title_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.6), Inches(1.2))
            title_paragraph = title_box.text_frame.paragraphs[0]
            title_paragraph.alignment = PP_ALIGN.CENTER
            title_run = title_paragraph.add_run()
            title_run.text = title
            title_run.font.size = Pt(38)
            title_run.font.bold = True
            title_run.font.color.rgb = title_color
            subtitle = " / ".join(str(point) for point in (info.get("points") or [])[:2])
            if subtitle:
                subtitle_box = slide.shapes.add_textbox(Inches(1.6), Inches(3.35), Inches(10.2), Inches(0.6))
                subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
                subtitle_paragraph.alignment = PP_ALIGN.CENTER
                subtitle_run = subtitle_paragraph.add_run()
                subtitle_run.text = subtitle
                subtitle_run.font.size = Pt(18)
                subtitle_run.font.color.rgb = body_color
        else:
            title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(11.9), Inches(0.8))
            title_run = title_box.text_frame.paragraphs[0].add_run()
            title_run.text = title
            title_run.font.size = Pt(28)
            title_run.font.bold = True
            title_run.font.color.rgb = title_color
            body_box = slide.shapes.add_textbox(Inches(0.95), Inches(1.65), Inches(11.0), Inches(4.6))
            body_frame = body_box.text_frame
            body_frame.word_wrap = True
            for point_index, point in enumerate((info.get("points") or [])[:7]):
                paragraph = body_frame.paragraphs[0] if point_index == 0 else body_frame.add_paragraph()
                paragraph.level = 0
                run = paragraph.add_run()
                run.text = f"• {point}"
                run.font.size = Pt(18)
                run.font.color.rgb = body_color

        footer = slide.shapes.add_textbox(Inches(0.75), Inches(6.9), Inches(11.8), Inches(0.3))
        footer_paragraph = footer.text_frame.paragraphs[0]
        footer_paragraph.alignment = PP_ALIGN.RIGHT
        footer_run = footer_paragraph.add_run()
        footer_run.text = f"{index + 1} / {len(slides)}"
        footer_run.font.size = Pt(10)
        footer_run.font.color.rgb = muted_color

    presentation.save(str(output_path))
    logger.info("[PPT image overlay] outline draft PPTX saved: %s", output_path)


def _move_shape_to_back(slide, shape) -> None:
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    insert_index = 2
    for index, child in enumerate(tree):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag in {"nvGrpSpPr", "grpSpPr"}:
            insert_index = index + 1
    tree.insert(insert_index, shape._element)
