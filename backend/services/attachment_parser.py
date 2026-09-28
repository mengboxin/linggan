"""Lightweight attachment text extraction for generation workflows."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, MutableMapping
from xml.etree import ElementTree as ET


MAX_EXTRACTED_CHARS = 30000


@dataclass
class ParsedAttachment:
    filename: str
    kind: str
    text: str
    size: int
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "filename": self.filename,
            "kind": self.kind,
            "text": self.text,
            "size": self.size,
            "warnings": self.warnings,
        }


def _limit_text(text: str, limit: int = MAX_EXTRACTED_CHARS) -> str:
    cleaned = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n").replace("\r", "\n")).strip()
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[:limit].rstrip() + "\n\n[内容已截断]"


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _xml_text_nodes(xml_bytes: bytes, tag_suffix: str) -> list[str]:
    root = ET.fromstring(xml_bytes)
    parts: list[str] = []
    for elem in root.iter():
        if elem.tag.endswith(tag_suffix) and elem.text:
            parts.append(elem.text)
    return parts


def _extract_docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    paragraphs: list[str] = []
    for para in root.iter():
        if not para.tag.endswith("}p"):
            continue
        parts = [node.text for node in para.iter() if node.tag.endswith("}t") and node.text]
        line = "".join(parts).strip()
        if line:
            paragraphs.append(line)
    return "\n".join(paragraphs)


def _slide_sort_key(name: str) -> int:
    match = re.search(r"slide(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def _extract_pptx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        slide_names = sorted(
            [n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")],
            key=_slide_sort_key,
        )
        sections: list[str] = []
        for idx, name in enumerate(slide_names, 1):
            parts = _xml_text_nodes(zf.read(name), "}t")
            clean_parts = [p.strip() for p in parts if p and p.strip()]
            if clean_parts:
                sections.append(f"第 {idx} 页：\n" + "\n".join(clean_parts))
        notes_names = sorted(
            [n for n in zf.namelist() if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")],
            key=_slide_sort_key,
        )
        for idx, name in enumerate(notes_names, 1):
            parts = _xml_text_nodes(zf.read(name), "}t")
            clean_parts = [p.strip() for p in parts if p and p.strip()]
            if clean_parts:
                sections.append(f"第 {idx} 页备注：\n" + "\n".join(clean_parts))
    return "\n\n".join(sections)


def _extract_pdf(data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(io.BytesIO(data))
        pages = []
        for idx, page in enumerate(reader.pages[:80], 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"第 {idx} 页：\n{text.strip()}")
        if len(reader.pages) > 80:
            warnings.append("PDF 页数较多，仅解析前 80 页")
        return "\n\n".join(pages), warnings
    except ImportError:
        warnings.append("当前环境未安装 pypdf，PDF 只能做有限文本兜底解析")
    except Exception as exc:
        warnings.append(f"PDF 解析失败，已尝试兜底：{exc}")

    # Best-effort fallback for simple text PDFs.
    decoded = _decode_text(data)
    fragments = re.findall(r"\(([^()]{2,200})\)", decoded)
    text = "\n".join(fragments)
    return text, warnings


def _extract_xlsx(data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    try:
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections: list[str] = []
        for sheet in wb.worksheets[:20]:
            rows: list[str] = []
            for ridx, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if ridx > 120:
                    warnings.append(f"{sheet.title} 仅解析前 120 行")
                    break
                values = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in values):
                    rows.append(",".join(values))
            if rows:
                sections.append(f"[{sheet.title}]\n" + "\n".join(rows))
        if len(wb.worksheets) > 20:
            warnings.append("工作表较多，仅解析前 20 个")
        return "\n\n".join(sections), warnings
    except ImportError:
        return "", ["当前环境未安装 openpyxl，暂不能解析 xlsx"]
    except Exception as exc:
        return "", [f"xlsx 解析失败：{exc}"]


def _extract_csv(data: bytes) -> str:
    raw = _decode_text(data)
    sample = raw[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except Exception:
        dialect = csv.excel
    rows: list[str] = []
    reader = csv.reader(io.StringIO(raw), dialect)
    for idx, row in enumerate(reader, 1):
        if idx > 160:
            rows.append("[CSV 已截断，仅解析前 160 行]")
            break
        rows.append(",".join(row))
    return "\n".join(rows)


def _kind_from_suffix(suffix: str) -> str:
    return {
        ".pptx": "pptx",
        ".ppt": "ppt",
        ".docx": "docx",
        ".doc": "doc",
        ".pdf": "pdf",
        ".xlsx": "xlsx",
        ".xls": "xls",
        ".csv": "csv",
        ".txt": "text",
        ".md": "markdown",
        ".json": "json",
    }.get(suffix.lower(), "file")


def parse_attachment_bytes(filename: str, data: bytes, content_type: str = "") -> ParsedAttachment:
    suffix = Path(filename).suffix.lower()
    kind = _kind_from_suffix(suffix)
    warnings: list[str] = []
    text = ""

    if suffix in {".txt", ".md"} or content_type.startswith("text/"):
        text = _decode_text(data)
    elif suffix == ".csv":
        text = _extract_csv(data)
    elif suffix == ".json":
        decoded = _decode_text(data)
        try:
            text = json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
        except Exception:
            text = decoded
    elif suffix == ".docx":
        try:
            text = _extract_docx(data)
        except Exception as exc:
            warnings.append(f"docx 解析失败：{exc}")
    elif suffix == ".pptx":
        try:
            text = _extract_pptx(data)
        except Exception as exc:
            warnings.append(f"pptx 解析失败：{exc}")
    elif suffix == ".pdf":
        text, warnings = _extract_pdf(data)
    elif suffix == ".xlsx":
        text, xlsx_warnings = _extract_xlsx(data)
        warnings.extend(xlsx_warnings)
    elif suffix in {".ppt", ".doc", ".xls"}:
        warnings.append("旧版二进制 Office 文件解析有限，请尽量另存为 pptx/docx/xlsx 后上传")
        text = ""
    else:
        text = _decode_text(data)
        if not text.strip():
            warnings.append("暂不支持该附件类型的自动解析")

    text = _limit_text(text)
    if not text:
        text = "[未提取到可用文本]"
    return ParsedAttachment(filename=filename, kind=kind, text=text, size=len(data), warnings=warnings)


def build_attachment_context(items: Iterable[dict | ParsedAttachment], limit: int = 60000) -> str:
    blocks: list[str] = []
    for idx, item in enumerate(items, 1):
        if isinstance(item, ParsedAttachment):
            filename = item.filename
            kind = item.kind
            text = item.text
        else:
            filename = str(item.get("filename") or f"附件 {idx}")
            kind = str(item.get("kind") or "file")
            text = str(item.get("text") or "")
        if not text.strip():
            continue
        blocks.append(f"附件 {idx}：{filename}（{kind}）\n{text.strip()}")
    return _limit_text("\n\n---\n\n".join(blocks), limit)


def _attachment_dict(item: dict | ParsedAttachment) -> dict | None:
    if isinstance(item, ParsedAttachment):
        return item.to_dict()
    if not isinstance(item, dict):
        return None
    filename = str(item.get("filename") or "").strip()
    if not filename:
        return None
    warnings = item.get("warnings")
    return {
        "filename": filename,
        "kind": str(item.get("kind") or "file"),
        "text": str(item.get("text") or ""),
        "size": int(item.get("size") or 0),
        "warnings": list(warnings) if isinstance(warnings, list) else [],
    }


def merge_attachments(
    existing: Iterable[dict | ParsedAttachment] | None,
    incoming: Iterable[dict | ParsedAttachment] | None,
    *,
    max_items: int = 16,
) -> list[dict]:
    """Merge parsed attachments, replacing older files with the same name."""
    merged: list[dict] = []
    positions: dict[str, int] = {}
    for raw_item in [*(existing or []), *(incoming or [])]:
        item = _attachment_dict(raw_item)
        if not item:
            continue
        key = item["filename"].casefold()
        previous = positions.get(key)
        if previous is None:
            positions[key] = len(merged)
            merged.append(item)
        else:
            merged[previous] = item
    if max_items <= 0:
        return []
    return merged[-max_items:]


def merge_attachment_state(
    state: MutableMapping[str, object],
    incoming: Iterable[dict | ParsedAttachment] | None,
    incoming_context: str = "",
    *,
    max_items: int = 16,
    context_limit: int = 60000,
) -> list[dict]:
    """Merge refinement attachments into a durable workflow state."""
    existing_items = state.get("attachments")
    current_items = existing_items if isinstance(existing_items, list) else []
    new_items = list(incoming or [])
    merged = merge_attachments(current_items, new_items, max_items=max_items)

    context_parts: list[str] = []
    rebuilt_context = build_attachment_context(merged, limit=context_limit)
    if rebuilt_context:
        context_parts.append(rebuilt_context)

    existing_context = str(state.get("attachment_context") or "").strip()
    if existing_context and not current_items:
        context_parts.append(existing_context)

    explicit_context = str(incoming_context or "").strip()
    if explicit_context and not new_items:
        context_parts.append(explicit_context)

    deduped_parts: list[str] = []
    for part in context_parts:
        if part and part not in deduped_parts:
            deduped_parts.append(part)

    state["attachments"] = merged
    state["attachment_context"] = _limit_text(
        "\n\n---\n\n".join(deduped_parts),
        context_limit,
    )
    return merged
